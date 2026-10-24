#!/usr/bin/env python3
"""Create a two-page Chinese advisor memo and editable cover letter."""
from pathlib import Path
import re,html
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches,Pt
D=Path(__file__).resolve().parents[1]
pdfmetrics.registerFont(TTFont('WQY','/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',subfontIndex=0))
body=ParagraphStyle('body',fontName='WQY',fontSize=10.2,leading=15,spaceAfter=7,wordWrap='CJK')
heading=ParagraphStyle('heading',parent=body,fontSize=12,leading=17,spaceBefore=7,spaceAfter=6,textColor=colors.HexColor('#19354b'))
title=ParagraphStyle('title',parent=heading,fontSize=16,leading=23,spaceAfter=14)
small=ParagraphStyle('small',parent=body,fontSize=9,leading=12,spaceAfter=2)
text=(D/'ADVISOR_REVIEW_SUMMARY.md').read_text();story=[];lines=text.splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('# '):story.append(Paragraph(html.escape(line[2:]),title));continue
 if line.startswith('## '):
  if line.startswith('## 5.'):story.append(PageBreak())
  story.append(Paragraph(html.escape(line[3:]),heading));continue
 if line.startswith('|'):
  rows=[]
  while True:
   vals=[v.strip() for v in line.strip('|').split('|')]
   if not all(re.fullmatch(r'[:\- ]+',v) for v in vals):rows.append([Paragraph(html.escape(v),small) for v in vals])
   if i>=len(lines) or not lines[i].strip().startswith('|'):break
   line=lines[i].strip();i+=1
  t=Table(rows,colWidths=[110,100,100,90,85],repeatRows=1)
  t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9eef2')),('LINEBELOW',(0,0),(-1,0),.5,colors.gray),('LINEBELOW',(0,-1),(-1,-1),.5,colors.gray),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
  story.extend([t,Spacer(1,8)]);continue
 story.append(Paragraph(html.escape(line),body))
def footer(c,doc):
 c.setFont('WQY',8);c.setFillColor(colors.gray);c.drawString(42,26,'导师审阅稿 · 尚未投稿 · 作者信息待确认');c.drawRightString(A4[0]-42,26,str(doc.page))
SimpleDocTemplate(str(D/'ADVISOR_REVIEW_SUMMARY.pdf'),pagesize=A4,rightMargin=42,leftMargin=42,topMargin=34,bottomMargin=42,title='R27 导师审阅摘要',author='Author information pending approval').build(story,onFirstPage=footer,onLaterPages=footer)
# Business-letter DOCX, preserving all explicit author-confirmation placeholders.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.85);sec.top_margin=sec.bottom_margin=Inches(.8)
style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(11);style.paragraph_format.space_after=Pt(6)
for line in (D/'COVER_LETTER_DRAFT.md').read_text().splitlines():
 if not line.strip():continue
 p=doc.add_paragraph()
 if line.startswith('# '):p.style='Title';p.add_run(line[2:])
 else:
  pieces=re.split(r'(\*\*.*?\*\*)',line)
  for part in pieces:
   run=p.add_run(part[2:-2] if part.startswith('**') else part);run.bold=part.startswith('**')
doc.save(D/'COVER_LETTER_DRAFT.docx')
highlights=[s[2:] for s in (D/'HIGHLIGHTS.md').read_text().splitlines() if s.startswith('- ')]
assert all(len(s)<=85 for s in highlights)
(D/'HIGHLIGHTS.txt').write_text('\n'.join(highlights)+'\n')
# Human-readable compiled bibliography, keeping IEEE ordering and a note on repeated authors.
s=(D/'REFERENCES_COMPILED.bbl').read_text();entries=re.split(r'\\bibitem\{([^}]+)\}',s)[1:];out=['# References\n','Canonical editable fields are in `REFERENCES.bib`; IEEE-formatted typeset references are in the paper. Repeated-author dashes follow IEEEtran style and denote the preceding authors. URLs/DOIs and version limits remain explicit.\n']
for n in range(0,len(entries),2):
 key,t=entries[n:n+2];t=t.split('\\end{thebibliography}')[0]
 t=re.sub(r'\\BIBentry[A-Za-z]+','',t);t=t.replace('\\emph{','{').replace('\\url{','{').replace('\\,','').replace('~',' ').replace('\\"O','Ö').replace('---','—').replace('--','–').replace('``','“').replace("''",'”')
 t=t.replace('{','').replace('}','');t=' '.join(t.split())
 out.append(f'{n//2+1}. **{key}** — {t}\n')
(D/'REFERENCES.md').write_text('\n'.join(out))
print('Advisor memo, cover letter, highlights and references generated')
