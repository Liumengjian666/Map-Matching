#!/usr/bin/env python3
"""Append reproducible public-dataset findings and figures to the phase Word report."""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import rosbag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from scipy.spatial.transform import Rotation, Slerp


DESKTOP = Path("/home/jian/桌面")
BASE_REPORT = DESKTOP / "先验地图定位论文创新阶段报告_2026-09-09.docx"
OUTPUT_REPORT = DESKTOP / "先验地图定位论文创新阶段报告_公开真值补充版_2026-09-09.docx"
ASSET_DIR = DESKTOP / "先验地图定位论文创新阶段报告_公开真值补充版_2026-09-09_assets"
RESULT_BAG = Path("/home/jian/rosbag/paper_localization/iilabs_loop_full/iilabs_loop_full_reliability_eskf_2x/result.bag")
GROUND_TRUTH = Path("/home/jian/rosbag/public_datasets/IILABS3D/iilabs3d_dataset/benchmark/livox_mid-360/loop/ground_truth.tum")


def load_truth():
    truth = np.loadtxt(GROUND_TRUTH)
    times = truth[:, 0]
    positions = truth[:, 1:4]
    rotations = Rotation.from_quat(truth[:, 4:8])
    offset = np.array([0.0, 0.0, 0.4612])
    sensor_positions = positions + rotations.apply(np.repeat(offset[None, :], len(truth), axis=0))
    return times, sensor_positions, rotations


def load_estimate(topic):
    stamps, positions, quaternions = [], [], []
    with rosbag.Bag(str(RESULT_BAG)) as bag:
        for _, msg, _ in bag.read_messages(topics=[topic]):
            p = msg.pose.pose.position
            q = msg.pose.pose.orientation
            stamps.append(msg.header.stamp.to_sec())
            positions.append((p.x, p.y, p.z))
            quaternions.append((q.x, q.y, q.z, q.w))
    return np.asarray(stamps), np.asarray(positions), Rotation.from_quat(quaternions)


def interpolate_truth(stamps, truth_times, sensor_positions, truth_rotations):
    right = np.searchsorted(truth_times, stamps).clip(1, len(truth_times) - 1)
    left = right - 1
    alpha = (stamps - truth_times[left]) / (truth_times[right] - truth_times[left])
    interpolated = ((1.0 - alpha[:, None]) * sensor_positions[left] +
                    alpha[:, None] * sensor_positions[right])
    initial_rotation = truth_rotations[0]
    gt_positions = initial_rotation.inv().apply(interpolated - sensor_positions[0])
    gt_rotations = initial_rotation.inv() * Slerp(truth_times, truth_rotations)(stamps)
    return gt_positions, gt_rotations


def make_figures():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    truth_times, sensor_positions, truth_rotations = load_truth()
    topics = {
        "NDT 10 Hz": "/dog_livo/ndt_odom",
        "Adaptive ESKF 10 Hz": "/dog_livo/odom_corrected",
        "High-rate 200 Hz": "/dog_livo/odom_high_rate",
    }
    series = {}
    for label, topic in topics.items():
        stamps, positions, rotations = load_estimate(topic)
        valid = (stamps >= truth_times[0]) & (stamps <= truth_times[-1])
        stamps, positions, rotations = stamps[valid], positions[valid], rotations[valid]
        gt_positions, gt_rotations = interpolate_truth(
            stamps, truth_times, sensor_positions, truth_rotations)
        series[label] = {
            "time": stamps - truth_times[0],
            "position": positions,
            "gt_position": gt_positions,
            "translation_error": np.linalg.norm(positions - gt_positions, axis=1),
            "rotation_error": (gt_rotations.inv() * rotations).magnitude() * 180.0 / np.pi,
        }

    fig, axes = plt.subplots(2, 1, figsize=(9.2, 5.8), sharex=True)
    for label, values in series.items():
        step = 20 if "200" in label else 1
        axes[0].plot(values["time"][::step], values["translation_error"][::step], label=label, linewidth=1.0)
        axes[1].plot(values["time"][::step], values["rotation_error"][::step], label=label, linewidth=1.0)
    axes[0].set_ylabel("Translation error (m)")
    axes[1].set_ylabel("Rotation error (deg)")
    axes[1].set_xlabel("Time from GT start (s)")
    axes[0].grid(alpha=0.25)
    axes[1].grid(alpha=0.25)
    axes[0].legend(ncol=3, fontsize=8)
    fig.tight_layout()
    error_path = ASSET_DIR / "iilabs_full_error_over_time.png"
    fig.savefig(error_path, dpi=180)
    plt.close(fig)

    ndt = series["NDT 10 Hz"]
    fig, ax = plt.subplots(figsize=(7.6, 5.7))
    ax.plot(ndt["gt_position"][:, 0], ndt["gt_position"][:, 1], label="Ground truth", linewidth=2.0)
    ax.plot(ndt["position"][:, 0], ndt["position"][:, 1], label="NDT estimate", linewidth=1.0)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel("X (m)")
    ax.set_ylabel("Y (m)")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    trajectory_path = ASSET_DIR / "iilabs_full_xy_trajectory.png"
    fig.savefig(trajectory_path, dpi=180)
    plt.close(fig)

    labels = ["Before fix\n10 Hz", "After fix\n10 Hz", "Before fix\n200 Hz", "After fix\n200 Hz"]
    rotation_rmse = [2.7509, 1.1193, 2.8654, 0.9843]
    fig, ax = plt.subplots(figsize=(7.2, 3.8))
    bars = ax.bar(labels, rotation_rmse, color=["#d95f5f", "#4c9f70", "#d95f5f", "#4c9f70"])
    ax.set_ylabel("30 s rotation RMSE (deg)")
    ax.grid(axis="y", alpha=0.25)
    for bar, value in zip(bars, rotation_rmse):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.05, f"{value:.3f}", ha="center", fontsize=9)
    fig.tight_layout()
    fix_path = ASSET_DIR / "gravity_correction_ablation.png"
    fig.savefig(fix_path, dpi=180)
    plt.close(fig)
    return error_path, trajectory_path, fix_path


def add_caption(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(9)


def build_report():
    error_path, trajectory_path, fix_path = make_figures()
    document = Document(BASE_REPORT)
    document.add_page_break()
    document.add_heading("公开真值数据集补充验证（IILABS3D）", level=1)
    document.add_paragraph(
        "本节补充 Livox Mid-360 公开数据的真值验证。测试序列为 IILABS3D indoor loop，"
        "原始 bag 共 623.57 s、6235 帧 PointCloud2 和约 200 Hz IMU。先验地图由同一序列每 5 帧取 1 帧，"
        "结合 ground_truth.tum 投影并以 0.15 m 体素降采样得到 534857 点。该设置可验证定位链路和重复结构鲁棒性，"
        "但存在同序列建图/定位的数据复用，不能替代跨序列、跨时段的独立测试。"
    )

    document.add_heading("1. 完整序列正式结果", level=2)
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["输出", "频率", "ATE平移RMSE", "ATE旋转RMSE", "RPE-1s平移RMSE", "RPE-1s旋转RMSE", "帧数"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    rows = [
        ("NDT", "10.000 Hz", "1.673 m", "15.778°", "0.2256 m", "2.529°", "6235"),
        ("自适应ESKF", "10.000 Hz", "1.670 m", "15.706°", "0.2256 m", "2.413°", "6235"),
        ("高频传播", "199.999 Hz", "1.665 m", "15.744°", "0.2257 m", "2.371°", "124513"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.add_paragraph(
        "完整序列 6235/6235 帧 NDT 收敛，融合接受 6235/6235、失败 0。纯定位计算耗时中位数 10.09 ms，"
        "P95 为 10.81 ms；2 倍速回放时 NDT 进程 CPU 中位数约 20.2%，ESKF 约 3.8%。"
        "但 RMSE 显著高于中位误差（NDT 平移中位数仅 0.205 m），表明误差不是均匀漂移，而是周期性灾难失效。"
    )
    document.add_picture(str(error_path), width=Inches(6.5))
    add_caption(document, "图1  完整序列无对齐真值误差随时间变化（周期性重复走廊失效清晰可见）")
    document.add_picture(str(trajectory_path), width=Inches(5.8))
    add_caption(document, "图2  IILABS3D 完整序列 XY 真值轨迹与 NDT 轨迹")

    document.add_heading("2. 已确认的姿态修复", level=2)
    document.add_paragraph(
        "原配置未显式限定连续重力方向校正，程序采用 1.5 g 的宽容差，会把运动加速度当作重力并污染 roll/pitch。"
        "同时，启动静止段只估计重力方向而未初始化陀螺零偏。创新工作区现已："
        "（1）关闭不安全的连续重力伪观测；（2）用前 200 帧静止 IMU 初始化陀螺零偏；"
        "（3）保留 10 Hz NDT 姿态观测和约 200 Hz 陀螺传播。提交号为 ccc0f15。"
    )
    document.add_picture(str(fix_path), width=Inches(5.8))
    add_caption(document, "图3  同一 30 秒公开片段修复前后姿态 RMSE")
    document.add_paragraph(
        "修复后，10 Hz 融合姿态 RMSE 由 2.751°降至 1.119°（下降约 59.3%），200 Hz 高频姿态由 2.865°降至 0.984°"
        "（下降约 65.7%）。修复后 30 秒片段中，高频平移 RMSE 为 0.169 m，相比 NDT 的 0.212 m 下降约 20.2%；"
        "高频 1 秒 RPE 为 0.0363 m/0.360°，优于 NDT 的 0.0383 m/0.485°。"
    )

    document.add_heading("3. 负结果与方法边界", level=2)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["150秒方案", "平移ATE", "旋转ATE", "平移RPE-1s", "结论"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    ablations = [
        ("原0.8m NDT基线", "1.519 m", "14.332°", "0.206 m", "参照"),
        ("0.4m细分辨率", "1.512 m", "14.733°", "0.238 m", "更慢且无稳定收益"),
        ("原始IMU旋转初值", "1.523 m", "14.246°", "0.202 m", "仅局部微小收益"),
        ("IMU+帧间ICP", "1.574 m", "14.337°", "0.191 m", "RPE改善但ATE恶化"),
        ("完整融合位姿反馈", "4.063 m", "26.985°", "0.206 m", "正反馈发散"),
        ("仅融合姿态反馈", "5.740 m", "51.541°", "0.325 m", "正反馈发散"),
    ]
    for row in ablations:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.add_paragraph(
        "上述失败方案均未提交到代码库，其 result.bag 和真值 JSON 保留在 /home/jian/rosbag/paper_localization/ 下作为负消融证据。"
        "结论是：单初值 NDT 与融合状态直接闭环会放大偏差；单纯提高分辨率或加入短时 ICP，只改善局部 RPE，"
        "无法解决地图中的全局结构别名。"
    )

    document.add_heading("4. 当前论文判断与下一步", level=2)
    document.add_paragraph(
        "当前版本已经具备可复现实验链路、真值评估、可靠度遥测和一项有明确量化收益的姿态修复，但尚不足以直接投稿。"
        "核心缺口已从笼统的“继续调参”收敛为两个可验证问题："
        "第一，设计多假设匹配与候选间歧义分数，显式识别重复走廊中的全局别名；"
        "第二，建立烟雾强度分级的点云退化协议，并使用独立建图/定位序列验证跨时段泛化。"
        "只有在独立真值测试上同时降低 ATE、尾部误差和失效率，且保持 Orin 实时性，才能形成三区论文的主要贡献。"
    )
    document.add_paragraph(
        "可复现路径：正式公开结果位于 /home/jian/rosbag/paper_localization/iilabs_loop_full/"
        "iilabs_loop_full_reliability_eskf_2x；公开数据和真值位于 /home/jian/rosbag/public_datasets/IILABS3D/；"
        "创新工作区为 /home/jian/livox_ws/dog_light_loc_paper_ws，当前分支 research/paper-innovation。"
    )

    for style_name in ("Normal", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Noto Sans CJK SC"
    document.save(OUTPUT_REPORT)
    print(OUTPUT_REPORT)


if __name__ == "__main__":
    build_report()
